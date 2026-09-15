#!/usr/bin/env python3
"""
Proposal Writing Service for Mr Bubba Services
Generates professional business proposals (PDF/Word format) from client input.

Usage:
    python3 generate_proposal.py --client "Client Name" --project "Project Title" \
        --scope "Project Scope" --timeline "Timeline" --budget "Budget" \
        --output-format pdf --email-to "client@example.com"
"""

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

SERVICE_EMAIL = os.environ.get("AGENTMAIL_EMAIL", "mrbubba@agentmail.to")
SERVICE_NAME = "Mr Bubba Services"
SERVICE_PHONE = "(555) 012-3456"
SERVICE_WEBSITE = "https://mrbubbaservices.com"


def generate_proposal_content(client_name: str, project_title: str, scope: str,
                               timeline: str, budget: str, notes: str = "") -> dict:
    """Generate structured proposal content from client inputs."""
    today = datetime.now().strftime("%B %d, %Y")
    proposal_id = f"HDS-{datetime.now().strftime('%Y%m%d')}-{abs(hash(client_name)) % 10000:04d}"
    
    content = {
        "proposal_id": proposal_id,
        "date": today,
        "client_name": client_name,
        "project_title": project_title,
        "scope": scope,
        "timeline": timeline,
        "budget": budget,
        "notes": notes,
        "executive_summary": (
            f"{SERVICE_NAME} is pleased to submit this proposal for the {project_title} "
            f"on behalf of {client_name}. Our team brings deep expertise in data analytics, "
            f"business intelligence, and custom data solutions tailored to drive measurable outcomes. "
            f"This proposal outlines our approach, timeline, and investment for delivering exceptional results."
        ),
        "approach": [
            "Discovery & Requirements Analysis — Deep-dive into your business objectives, data landscape, and success criteria",
            "Solution Architecture — Design a scalable, maintainable solution aligned with your technical environment",
            "Implementation & Development — Build, test, and iterate on deliverables with regular stakeholder updates",
            "Quality Assurance & Validation — Rigorous testing to ensure accuracy, performance, and reliability",
            "Delivery & Knowledge Transfer — Deploy final deliverables with documentation and training",
            "Ongoing Support — Post-delivery support to ensure continued success and optimization"
        ],
        "deliverables": [
            "Complete project deliverables as defined in scope",
            "Technical documentation and user guides",
            "Data models and architectural diagrams",
            "Source code and configuration files (where applicable)",
            "Training sessions and knowledge transfer materials"
        ],
        "terms": [
            f"Total project investment: {budget}",
            f"Estimated timeline: {timeline}",
            "50% deposit due upon acceptance, 50% upon final delivery",
            "Additional scope changes will be billed at $150/hour",
            "All work is confidential and intellectual property transfers upon full payment"
        ]
    }
    return content


def generate_word_proposal(content: dict, output_path: str) -> str:
    """Generate a Word (.docx) proposal document."""
    if not HAS_DOCX:
        return ""
    
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Title
    title = doc.add_heading(f'{SERVICE_NAME}', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtitle
    subtitle = doc.add_heading('Business Proposal', level=1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Proposal info table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.style = 'Light Grid Accent 1'
    info_data = [
        ("Proposal ID:", content["proposal_id"]),
        ("Date:", content["date"]),
        ("Prepared For:", content["client_name"]),
        ("Project:", content["project_title"]),
    ]
    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
    
    doc.add_paragraph()  # spacer
    
    # Executive Summary
    doc.add_heading('Executive Summary', level=2)
    doc.add_paragraph(content["executive_summary"])
    
    # Scope of Work
    doc.add_heading('Scope of Work', level=2)
    doc.add_paragraph(content["scope"])
    
    # Approach
    doc.add_heading('Our Approach', level=2)
    for step in content["approach"]:
        doc.add_paragraph(step, style='List Number')
    
    # Deliverables
    doc.add_heading('Deliverables', level=2)
    for item in content["deliverables"]:
        doc.add_paragraph(item, style='List Bullet')
    
    # Timeline
    doc.add_heading('Timeline', level=2)
    doc.add_paragraph(content["timeline"])
    
    # Investment
    doc.add_heading('Investment', level=2)
    doc.add_paragraph(content["budget"])
    
    # Terms
    doc.add_heading('Terms & Conditions', level=2)
    for term in content["terms"]:
        doc.add_paragraph(term, style='List Bullet')
    
    # Closing
    doc.add_paragraph()
    doc.add_paragraph(f"We look forward to the opportunity to partner with {content['client_name']} on this exciting initiative.")
    doc.add_paragraph()
    doc.add_paragraph(f"Sincerely,")
    doc.add_paragraph()
    doc.add_paragraph(f"The {SERVICE_NAME} Team")
    doc.add_paragraph(f"Email: {SERVICE_EMAIL}")
    doc.add_paragraph(f"Phone: {SERVICE_PHONE}")
    
    doc.save(output_path)
    return output_path


def generate_pdf_proposal(content: dict, output_path: str) -> str:
    """Generate a PDF proposal document."""
    if not HAS_REPORTLAB:
        return ""
    
    doc = SimpleDocTemplate(output_path, pagesize=letter,
                            rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle('CustomTitle', parent=styles['Title'],
                                  fontSize=24, spaceAfter=20, alignment=1)
    heading_style = ParagraphStyle('CustomHeading', parent=styles['Heading1'],
                                    fontSize=16, spaceAfter=12, spaceBefore=12)
    body_style = ParagraphStyle('CustomBody', parent=styles['Normal'],
                                 fontSize=11, spaceAfter=10, leading=14)
    
    story = []
    
    # Title
    story.append(Paragraph(SERVICE_NAME, title_style))
    story.append(Paragraph("Business Proposal", title_style))
    story.append(Spacer(1, 20))
    
    # Info section
    info_data = [
        f"<b>Proposal ID:</b> {content['proposal_id']}",
        f"<b>Date:</b> {content['date']}",
        f"<b>Prepared For:</b> {content['client_name']}",
        f"<b>Project:</b> {content['project_title']}",
    ]
    for info in info_data:
        story.append(Paragraph(info, body_style))
    
    story.append(Spacer(1, 20))
    
    # Executive Summary
    story.append(Paragraph("Executive Summary", heading_style))
    story.append(Paragraph(content["executive_summary"], body_style))
    
    # Scope
    story.append(Paragraph("Scope of Work", heading_style))
    story.append(Paragraph(content["scope"], body_style))
    
    # Approach
    story.append(Paragraph("Our Approach", heading_style))
    for i, step in enumerate(content["approach"], 1):
        story.append(Paragraph(f"{i}. {step}", body_style))
    
    # Deliverables
    story.append(Paragraph("Deliverables", heading_style))
    for item in content["deliverables"]:
        story.append(Paragraph(f"• {item}", body_style))
    
    # Timeline
    story.append(Paragraph("Timeline", heading_style))
    story.append(Paragraph(content["timeline"], body_style))
    
    # Investment
    story.append(Paragraph("Investment", heading_style))
    story.append(Paragraph(content["budget"], body_style))
    
    # Terms
    story.append(Paragraph("Terms & Conditions", heading_style))
    for term in content["terms"]:
        story.append(Paragraph(f"• {term}", body_style))
    
    story.append(Spacer(1, 30))
    story.append(Paragraph(f"We look forward to the opportunity to partner with {content['client_name']} on this exciting initiative.", body_style))
    story.append(Spacer(1, 20))
    story.append(Paragraph("Sincerely,", body_style))
    story.append(Paragraph(f"The {SERVICE_NAME} Team", body_style))
    story.append(Paragraph(f"Email: {SERVICE_EMAIL}", body_style))
    story.append(Paragraph(f"Phone: {SERVICE_PHONE}", body_style))
    
    doc.build(story)
    return output_path


def generate_proposal(client_name: str, project_title: str, scope: str,
                      timeline: str, budget: str, notes: str = "",
                      output_format: str = "pdf", output_dir: str = ".") -> dict:
    """Main function to generate a proposal."""
    content = generate_proposal_content(client_name, project_title, scope, timeline, budget, notes)
    
    os.makedirs(output_dir, exist_ok=True)
    
    result = {
        "content": content,
        "files": [],
        "status": "success"
    }
    
    if output_format in ("pdf", "both"):
        pdf_path = os.path.join(output_dir, f"{content['proposal_id']}.pdf")
        if generate_pdf_proposal(content, pdf_path):
            result["files"].append(pdf_path)
    
    if output_format in ("docx", "both"):
        docx_path = os.path.join(output_dir, f"{content['proposal_id']}.docx")
        if generate_word_proposal(content, docx_path):
            result["files"].append(docx_path)
    
    if not result["files"]:
        result["status"] = "partial"
        result["warning"] = "No PDF/Word libraries available; content returned in JSON only"
    
    return result


def main():
    parser = argparse.ArgumentParser(description="Generate a professional business proposal")
    parser.add_argument("--client", required=True, help="Client/company name")
    parser.add_argument("--project", required=True, help="Project title")
    parser.add_argument("--scope", required=True, help="Project scope description")
    parser.add_argument("--timeline", required=True, help="Estimated timeline")
    parser.add_argument("--budget", required=True, help="Project budget/investment")
    parser.add_argument("--notes", default="", help="Additional notes")
    parser.add_argument("--output-format", choices=["pdf", "docx", "both"], default="pdf",
                        help="Output format (default: pdf)")
    parser.add_argument("--output-dir", default="/home/vboxuser/mr-bubba-mission/services/proposals/output",
                        help="Output directory")
    parser.add_argument("--email-to", help="Client email for delivery")
    parser.add_argument("--json-only", action="store_true", help="Only output JSON content")
    
    args = parser.parse_args()
    
    if args.json_only:
        content = generate_proposal_content(args.client, args.project, args.scope,
                                            args.timeline, args.budget, args.notes)
        print(json.dumps(content, indent=2))
        return
    
    result = generate_proposal(args.client, args.project, args.scope, args.timeline,
                               args.budget, args.notes, args.output_format, args.output_dir)
    
    if args.email_to:
        result["email_to"] = args.email_to
        result["email_sent"] = "Use AgentMail API to deliver"
    
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
